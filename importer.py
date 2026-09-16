from io import BytesIO

from openpyxl import Workbook, load_workbook

from tools import TOOLS

EXPECTED_HEADER = ["portal", "username", "password"]


def resolve_tool_key(value):
    cleaned = value.strip().lower()

    for key, tool in TOOLS.items():
        if cleaned == key.lower() or cleaned == tool["name"].lower():
            return key

    return None


def build_template():
    wb = Workbook()
    ws = wb.active
    ws.title = "Credentials"
    ws.append(["Portal", "Username", "Password"])

    for tool in TOOLS.values():
        ws.append([tool["name"], "", ""])

    ws.freeze_panes = "A2"
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 28

    stream = BytesIO()
    wb.save(stream)
    return stream.getvalue()


def parse_import(fileobj):
    try:
        wb = load_workbook(fileobj, read_only=True, data_only=True)
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
    except Exception:
        raise ValueError("Could not read the file. Please upload a valid .xlsx file.")

    header = next(rows, None)

    if not header or _normalize_header(header) != EXPECTED_HEADER:
        raise ValueError(
            "Invalid format. Columns must be: Portal, Username, Password"
        )

    created = []
    errors = []

    for idx, row in enumerate(rows, start=2):
        if row is None or all(
            cell is None or str(cell).strip() == "" for cell in row
        ):
            continue

        portal = _text(row[0])
        username = _text(row[1])
        password = _text(row[2])

        if not portal or not username or not password:
            errors.append((idx, "missing value"))
            continue

        tool_key = resolve_tool_key(portal)

        if not tool_key:
            errors.append((idx, f"unknown portal '{portal}'"))
            continue

        created.append((tool_key, username, password))

    wb.close()

    return created, errors


def _text(value):
    return str(value).strip() if value is not None else ""


def _normalize_header(header):
    return [_text(cell).lower() for cell in header[:3]]